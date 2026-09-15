"""Turns account report spreadsheets (and user-uploaded files) into compact text summaries
that can be fed to Gemini as context, without blowing past a reasonable prompt size."""

import csv
import json
from pathlib import Path

import openpyxl

MAX_SUMMARY_CHARS = 12000
MAX_ROWS_PER_SHEET = 25
MAX_COLS_PER_ROW = 12
MAX_CELL_CHARS = 60


def _format_cell(value: object) -> str:
    if value is None:
        return ""
    text = str(value)
    return text if len(text) <= MAX_CELL_CHARS else f"{text[:MAX_CELL_CHARS]}…"


def summarize_xlsx(path: Path, max_chars: int = MAX_SUMMARY_CHARS) -> str:
    """Summarize every worksheet in an .xlsx workbook: header row + a handful of sample rows."""
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    lines: list[str] = [f"Workbook: {path.name}", f"Sheets: {', '.join(workbook.sheetnames)}"]

    for sheet in workbook.worksheets:
        lines.append(f"\n--- Sheet: {sheet.title} (rows={sheet.max_row}, cols={sheet.max_column}) ---")
        row_count = 0
        for row in sheet.iter_rows(max_row=MAX_ROWS_PER_SHEET + 1, max_col=MAX_COLS_PER_ROW, values_only=True):
            cells = [_format_cell(value) for value in row]
            if not any(cells):
                continue
            lines.append(" | ".join(cells))
            row_count += 1
            if row_count > MAX_ROWS_PER_SHEET:
                lines.append("... (additional rows truncated)")
                break

    workbook.close()
    summary = "\n".join(lines)
    if len(summary) > max_chars:
        summary = summary[:max_chars] + "\n... (summary truncated)"
    return summary


def summarize_csv(path: Path, max_rows: int = 60, max_chars: int = MAX_SUMMARY_CHARS) -> str:
    lines: list[str] = [f"CSV file: {path.name}"]
    with path.open("r", newline="", encoding="utf-8", errors="replace") as csv_file:
        reader = csv.reader(csv_file)
        for index, row in enumerate(reader):
            if index >= max_rows:
                lines.append("... (additional rows truncated)")
                break
            lines.append(" | ".join(_format_cell(value) for value in row))
    summary = "\n".join(lines)
    if len(summary) > max_chars:
        summary = summary[:max_chars] + "\n... (summary truncated)"
    return summary


def summarize_json(path: Path, max_chars: int = MAX_SUMMARY_CHARS) -> str:
    payload = json.loads(path.read_text(encoding="utf-8", errors="replace"))
    text = json.dumps(payload, indent=2)[:max_chars]
    return f"JSON file: {path.name}\n{text}"


def summarize_text(path: Path, max_chars: int = MAX_SUMMARY_CHARS) -> str:
    text = path.read_text(encoding="utf-8", errors="replace")[:max_chars]
    return f"Text file: {path.name}\n{text}"


def summarize_file(path: Path, filename: str | None = None) -> str:
    """Dispatch to a format-specific summarizer based on file extension."""
    name = filename or path.name
    suffix = Path(name).suffix.lower()

    if suffix in {".xlsx", ".xlsm"}:
        return summarize_xlsx(path)
    if suffix == ".csv":
        return summarize_csv(path)
    if suffix == ".json":
        return summarize_json(path)
    if suffix in {".txt", ".md", ".log"}:
        return summarize_text(path)

    raise ValueError(f"Unsupported file type for chat bot upload: {suffix or '(no extension)'}")
